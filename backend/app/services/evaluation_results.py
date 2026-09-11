"""Leaderboard ranking: 70% raw percentage + 30% within-track percentile.

Raw scores are never overwritten. All derived statistics are computed on each
request so the leaderboard cannot go stale. Evaluation documents are not modified.

    RawPercentage = (raw_score / max_score) * 100
    TrackPercentile = 100 * (teams in this track with a strictly lower raw_score) / (n - 1)
    FinalScore = (RawPercentage * 0.70) + (TrackPercentile * 0.30)

Percentile is computed separately per track from completed evaluations only.
A single evaluated team in a track receives TrackPercentile = 100.
Z-scores are not used for ranking.
"""

from __future__ import annotations

from datetime import datetime, timezone
from io import BytesIO
from math import isfinite, sqrt
from typing import Any, Iterable, Optional

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

RAW_WEIGHT = 0.70
PERCENTILE_WEIGHT = 0.30
NORMALIZED_MIN = 0.0
NORMALIZED_MAX = 100.0

NORMALIZATION_EXPLANATION = (
    "Final score is 70% of the team's raw percentage plus 30% of its percentile "
    "among completed evaluations in the same track. Raw scores are unchanged."
)

STATUS_OK = "ok"
STATUS_ZERO_VARIANCE = "zero_variance"
STATUS_SMALL_SAMPLE = "unreliable_small_sample"
STATUS_NO_DATA = "no_evaluated_teams"


def round2(value: Optional[float]) -> Optional[float]:
    if value is None:
        return None
    return round(float(value), 2)


def clamp(value: float, lo: float = NORMALIZED_MIN, hi: float = NORMALIZED_MAX) -> float:
    return max(lo, min(hi, value))


def as_number(value: Any) -> Optional[float]:
    if isinstance(value, bool) or value is None:
        return None
    if isinstance(value, (int, float)):
        number = float(value)
        return number if isfinite(number) else None
    if isinstance(value, str):
        try:
            number = float(value.strip())
        except ValueError:
            return None
        return number if isfinite(number) else None
    return None


def mean(values: list[float]) -> float:
    if not values:
        return 0.0
    return sum(values) / len(values)


def population_stddev(values: list[float]) -> float:
    """Population standard deviation (divide by n). Zero when n < 2. Descriptive only — not used for ranking."""
    n = len(values)
    if n < 2:
        return 0.0
    avg = mean(values)
    variance = sum((value - avg) ** 2 for value in values) / n
    return sqrt(variance)


def track_percentile(raw_score: float, track_scores: list[float]) -> float:
    """Percentile rank of raw_score within one track, on a 0–100 scale.

    TrackPercentile = 100 * (count of strictly lower scores) / (n - 1)
    The unique highest score is 100. The unique lowest score is 0.
    A track with one evaluated team is treated as the 100th percentile.
    Unevaluated teams are never included in track_scores.
    """
    n = len(track_scores)
    if n <= 1:
        return 100.0
    below = sum(1 for score in track_scores if score < raw_score)
    return 100.0 * below / (n - 1)


def compute_final_score(raw_percentage: float, percentile: float) -> float:
    """FinalScore = (RawPercentage * 0.70) + (TrackPercentile * 0.30)."""
    return clamp(raw_percentage * RAW_WEIGHT + percentile * PERCENTILE_WEIGHT)


def _eval_timestamp(evaluation: dict) -> datetime:
    for key in ("updated_at", "submitted_at"):
        value = evaluation.get(key)
        if isinstance(value, datetime):
            return value if value.tzinfo else value.replace(tzinfo=timezone.utc)
    return datetime.min.replace(tzinfo=timezone.utc)


def _dedupe_judge_evaluations(evaluations: list[dict]) -> list[dict]:
    """Keep the latest evaluation per judge for a team (unique index already exists)."""
    latest: dict[str, dict] = {}
    for evaluation in evaluations:
        judge_id = evaluation.get("judge_id") or ""
        previous = latest.get(judge_id)
        if previous is None or _eval_timestamp(evaluation) >= _eval_timestamp(previous):
            latest[judge_id] = evaluation
    return list(latest.values())


def validate_evaluation(evaluation: dict, criteria_by_id: dict[str, dict]) -> tuple[Optional[float], dict[str, float], list[str]]:
    """Return (raw_total, cleaned_scores, issues). raw_total is None when the evaluation is unusable."""
    issues: list[str] = []
    if evaluation.get("status") != "submitted":
        issues.append("not_submitted")
        return None, {}, issues
    raw_scores = evaluation.get("scores")
    if not isinstance(raw_scores, dict) or not raw_scores:
        issues.append("missing_scores")
        return None, {}, issues

    cleaned: dict[str, float] = {}
    for criterion_id, criterion in criteria_by_id.items():
        value = as_number(raw_scores.get(criterion_id))
        max_marks = as_number(criterion.get("max_marks")) or 0
        if value is None:
            issues.append(f"non_numeric:{criterion_id}")
            continue
        if value < 0 or value > max_marks:
            issues.append(f"out_of_range:{criterion_id}")
            continue
        cleaned[criterion_id] = value

    extra = set(raw_scores) - set(criteria_by_id)
    if extra:
        issues.append("unknown_criteria")

    if len(cleaned) != len(criteria_by_id):
        issues.append("incomplete_criteria")
        return None, cleaned, issues
    return sum(cleaned.values()), cleaned, issues


def _sort_key(row: dict) -> tuple:
    return (
        -float(row["final_score"]),
        -float(row["raw_score"]),
        -float(row.get("track_percentile") or 0),
        str(row.get("team_name") or "").casefold(),
        str(row.get("registration_id") or "").casefold(),
    )


def assign_ranks(rows: list[dict], rank_key: str) -> None:
    """Dense sequential ranks after the deterministic sort. Ties are broken, not shared."""
    for index, row in enumerate(rows, start=1):
        row[rank_key] = index


def _public_team(registration: dict) -> dict:
    team = registration.get("team") or {}
    registration_id = registration.get("registration_id", "")
    return {
        "registration_id": registration_id,
        "reference_id": registration_id,
        "team_name": team.get("teamName", "") or "",
        "ps_id": team.get("psId", "") or "",
        "problem_statement": team.get("psTitle", "") or "",
        "theme": team.get("theme", "") or "",
        "domain": team.get("category", "") or "",
    }


def compute_results(
    *,
    tracks: list[dict],
    judges: list[dict],
    criteria: list[dict],
    evaluations: list[dict],
    registrations: list[dict],
    assigned_by_track: dict[str, list[str]],
) -> dict:
    """Build leaderboard rows, track statistics, and per-team evaluation details."""
    criteria_by_id = {c["id"]: c for c in criteria if c.get("is_active") and c.get("id")}
    criteria_max = sum(as_number(c.get("max_marks")) or 0 for c in criteria_by_id.values())
    ordered_criteria = sorted(criteria_by_id.values(), key=lambda c: (c.get("order", 0), c.get("name", "")))

    track_docs = {t["track_id"]: t for t in tracks if t.get("track_id")}
    judge_docs = {j.get("id") or j.get("judge_id"): j for j in judges}
    registrations_by_id = {r["registration_id"]: r for r in registrations if r.get("registration_id")}

    grouped: dict[tuple[str, str], list[dict]] = {}
    malformed: list[dict] = []
    for evaluation in evaluations:
        track_id = evaluation.get("track_id")
        registration_id = evaluation.get("registration_id")
        if not track_id or not registration_id:
            malformed.append(
                {
                    "evaluation_id": evaluation.get("evaluation_id"),
                    "reason": "missing_track_or_registration",
                }
            )
            continue
        total, cleaned, issues = validate_evaluation(evaluation, criteria_by_id)
        if total is None:
            malformed.append(
                {
                    "evaluation_id": evaluation.get("evaluation_id"),
                    "registration_id": registration_id,
                    "track_id": track_id,
                    "judge_id": evaluation.get("judge_id"),
                    "reason": ", ".join(issues) or "malformed",
                }
            )
            continue
        grouped.setdefault((track_id, registration_id), []).append({**evaluation, "_cleaned_scores": cleaned, "_raw_total": total})

    completed_rows: list[dict] = []
    details: dict[tuple[str, str], dict] = {}
    evaluated_ids_by_track: dict[str, set[str]] = {track_id: set() for track_id in track_docs}

    for (track_id, registration_id), raw_evals in grouped.items():
        track = track_docs.get(track_id)
        registration = registrations_by_id.get(registration_id)
        if not track or not registration:
            malformed.append(
                {
                    "registration_id": registration_id,
                    "track_id": track_id,
                    "reason": "missing_track_or_registration_record",
                }
            )
            continue

        unique_evals = _dedupe_judge_evaluations(raw_evals)
        judges_required = int(track.get("judges_required") or 1)
        if len(unique_evals) < judges_required:
            continue

        unique_evals.sort(key=_eval_timestamp)
        raw_score = mean([item["_raw_total"] for item in unique_evals])
        criterion_scores: dict[str, float] = {}
        for criterion_id in criteria_by_id:
            criterion_scores[criterion_id] = mean([item["_cleaned_scores"][criterion_id] for item in unique_evals])
        max_score = criteria_max
        raw_percentage = (raw_score / max_score * 100) if max_score else 0.0
        team = _public_team(registration)
        judge_ids = [item.get("judge_id") or "" for item in unique_evals]
        judge_names = [item.get("judge_name") or (judge_docs.get(jid) or {}).get("name", "") for jid, item in zip(judge_ids, unique_evals)]
        row = {
            **team,
            "track_id": track_id,
            "track_name": track.get("name", ""),
            "track_code": track.get("code", ""),
            "judge_id": judge_ids[0] if len(judge_ids) == 1 else ",".join(judge_ids),
            "judge_name": judge_names[0] if len(judge_names) == 1 else ", ".join(n for n in judge_names if n),
            "judge_ids": judge_ids,
            "judge_names": judge_names,
            "raw_score": raw_score,
            "score": raw_score,
            "max_score": max_score,
            "raw_percentage": raw_percentage,
            "judges_count": len(unique_evals),
            "judges_required": judges_required,
            "evaluation_status": "completed",
            "criterion_scores": criterion_scores,
            "evaluations": unique_evals,
        }
        completed_rows.append(row)
        evaluated_ids_by_track.setdefault(track_id, set()).add(registration_id)
        details[(track_id, registration_id)] = row

    by_track: dict[str, list[dict]] = {}
    for row in completed_rows:
        by_track.setdefault(row["track_id"], []).append(row)

    track_stats: list[dict] = []
    for track_id, track in track_docs.items():
        track_rows = by_track.get(track_id, [])
        raw_scores = [row["raw_score"] for row in track_rows]
        sample_size = len(raw_scores)
        track_mean = mean(raw_scores) if sample_size else None
        track_stddev = population_stddev(raw_scores) if sample_size else None
        assigned = list(dict.fromkeys(assigned_by_track.get(track_id, [])))
        evaluated_ids = evaluated_ids_by_track.get(track_id, set())
        assigned_set = set(assigned) | evaluated_ids
        pending_ids = [rid for rid in assigned if rid not in evaluated_ids]
        if sample_size == 0:
            status = STATUS_NO_DATA
        else:
            status = STATUS_OK

        track_judge_ids = [j.get("id") or j.get("judge_id") for j in judges if j.get("track_id") == track_id and j.get("is_active", True)]
        track_judge_names = [
            (judge_docs.get(jid) or {}).get("name", "")
            for jid in track_judge_ids
            if jid
        ]

        for row in track_rows:
            percentile = track_percentile(row["raw_score"], raw_scores)
            score = compute_final_score(row["raw_percentage"], percentile)
            row.update(
                {
                    "track_mean": track_mean,
                    "track_stddev": track_stddev,
                    "track_percentile": percentile,
                    "final_score": score,
                    "normalized_score": score,
                    "normalization_status": status,
                    "normalization_reliable": True,
                    "evaluated_in_track": sample_size,
                }
            )
        track_rows.sort(key=_sort_key)
        assign_ranks(track_rows, "track_rank")

        track_stats.append(
            {
                "track_id": track_id,
                "track_name": track.get("name", ""),
                "track_code": track.get("code", ""),
                "judge_ids": [jid for jid in track_judge_ids if jid],
                "judge_names": [name for name in track_judge_names if name],
                "total_teams": len(assigned_set),
                "evaluated_teams": sample_size,
                "pending_teams": len(pending_ids),
                "pending_registration_ids": pending_ids,
                "mean_raw_score": track_mean,
                "stddev": track_stddev,
                "minimum": min(raw_scores) if raw_scores else None,
                "maximum": max(raw_scores) if raw_scores else None,
                "normalization_status": status,
                "normalization_reliable": status == STATUS_OK,
            }
        )

    completed_rows.sort(key=_sort_key)
    assign_ranks(completed_rows, "overall_rank")
    for row in completed_rows:
        row["rank"] = row["overall_rank"]

    pending_rows: list[dict] = []
    pending_seen: set[tuple[str, str]] = set()
    for track_id, track in track_docs.items():
        evaluated_ids = evaluated_ids_by_track.get(track_id, set())
        for registration_id in assigned_by_track.get(track_id, []):
            key = (track_id, registration_id)
            if registration_id in evaluated_ids or key in pending_seen:
                continue
            registration = registrations_by_id.get(registration_id)
            if not registration:
                continue
            pending_seen.add(key)
            team = _public_team(registration)
            pending_rows.append(
                {
                    **team,
                    "track_id": track_id,
                    "track_name": track.get("name", ""),
                    "track_code": track.get("code", ""),
                    "judge_id": "",
                    "judge_name": "",
                    "raw_score": None,
                    "score": None,
                    "max_score": criteria_max,
                    "raw_percentage": None,
                    "track_mean": None,
                    "track_stddev": None,
                    "track_percentile": None,
                    "final_score": None,
                    "normalized_score": None,
                    "track_rank": None,
                    "overall_rank": None,
                    "rank": None,
                    "judges_count": 0,
                    "judges_required": int(track.get("judges_required") or 1),
                    "evaluation_status": "pending",
                    "normalization_status": STATUS_NO_DATA,
                    "normalization_reliable": False,
                }
            )

    assigned_ids = {rid for ids in assigned_by_track.values() for rid in ids}
    evaluated_ids = {row["registration_id"] for row in completed_rows}
    summary = {
        "tracks": len(track_docs),
        "total_teams": len(assigned_ids),
        "evaluated": len(evaluated_ids),
        "pending": len(assigned_ids - evaluated_ids),
        "malformed_evaluations": len(malformed),
        "criteria_max": criteria_max,
        "raw_weight": RAW_WEIGHT,
        "percentile_weight": PERCENTILE_WEIGHT,
    }

    return {
        "rows": completed_rows,
        "pending_rows": pending_rows,
        "track_stats": sorted(track_stats, key=lambda t: (t["track_name"] or "").casefold()),
        "malformed": malformed,
        "summary": summary,
        "criteria": ordered_criteria,
        "details": details,
        "criteria_max": criteria_max,
    }


def serialize_leaderboard_row(row: dict, finalists: Optional[int] = None) -> dict:
    overall_rank = row.get("overall_rank")
    is_finalist = bool(finalists and overall_rank and overall_rank <= finalists and row.get("evaluation_status") == "completed")
    return {
        "registration_id": row.get("registration_id", ""),
        "reference_id": row.get("reference_id") or row.get("registration_id", ""),
        "team_name": row.get("team_name", ""),
        "problem_statement": row.get("problem_statement", ""),
        "ps_id": row.get("ps_id", ""),
        "theme": row.get("theme", ""),
        "domain": row.get("domain", ""),
        "track_id": row.get("track_id", ""),
        "track_name": row.get("track_name", ""),
        "track_code": row.get("track_code", ""),
        "judge_id": row.get("judge_id", ""),
        "judge_name": row.get("judge_name", ""),
        "raw_score": round2(row.get("raw_score")),
        "score": round2(row.get("raw_score")),
        "max_score": round2(row.get("max_score")) or 0,
        "raw_percentage": round2(row.get("raw_percentage")),
        "track_mean": round2(row.get("track_mean")),
        "track_stddev": round2(row.get("track_stddev")),
        "track_percentile": round2(row.get("track_percentile")),
        "final_score": round2(row.get("final_score")),
        "normalized_score": round2(row.get("final_score") if row.get("final_score") is not None else row.get("normalized_score")),
        "track_rank": row.get("track_rank"),
        "overall_rank": overall_rank,
        "rank": overall_rank,
        "judges_count": row.get("judges_count", 0),
        "judges_required": row.get("judges_required", 1),
        "evaluation_status": row.get("evaluation_status", "pending"),
        "normalization_status": row.get("normalization_status"),
        "normalization_reliable": bool(row.get("normalization_reliable")),
        "is_finalist": is_finalist,
    }


def serialize_track_stat(stat: dict) -> dict:
    return {
        "track_id": stat["track_id"],
        "track_name": stat["track_name"],
        "track_code": stat.get("track_code", ""),
        "judge_ids": stat.get("judge_ids", []),
        "judge_names": stat.get("judge_names", []),
        "total_teams": stat["total_teams"],
        "evaluated_teams": stat["evaluated_teams"],
        "pending_teams": stat["pending_teams"],
        "mean_raw_score": round2(stat.get("mean_raw_score")),
        "stddev": round2(stat.get("stddev")),
        "minimum": round2(stat.get("minimum")),
        "maximum": round2(stat.get("maximum")),
        "normalization_status": stat["normalization_status"],
        "normalization_reliable": stat["normalization_reliable"],
    }


def serialize_team_result(row: dict, criteria: list[dict], finalists: Optional[int] = None) -> dict:
    payload = serialize_leaderboard_row(row, finalists)
    scores = row.get("criterion_scores") or {}
    payload["criteria"] = [
        {
            "criterion_id": criterion.get("id"),
            "name": criterion.get("name", ""),
            "score": round2(scores.get(criterion.get("id"))),
            "max_marks": criterion.get("max_marks"),
            "description": criterion.get("description", ""),
        }
        for criterion in criteria
    ]
    payload["normalization_explanation"] = NORMALIZATION_EXPLANATION
    payload["methodology"] = methodology_text()
    return payload


def methodology_text() -> dict:
    return {
        "raw_score": "Sum of the team's criterion scores. Criteria are unweighted; each uses its configured max_marks. Historical evaluation documents are never rewritten.",
        "raw_percentage": "RawPercentage = (raw_score / max_score) * 100. max_score is the sum of active criterion max_marks (currently 70).",
        "track_percentile": "TrackPercentile is computed separately for each track from completed evaluations only. TrackPercentile = 100 × (number of teams in the track with a strictly lower raw_score) / (n − 1). The unique highest score is 100; the unique lowest is 0. A track with one evaluated team is 100. Unevaluated teams are excluded.",
        "final_score": "FinalScore = (RawPercentage × 0.70) + (TrackPercentile × 0.30).",
        "normalized_score": "Alias of FinalScore. Z-scores are not used.",
        "ties": "Sort by FinalScore descending, then RawScore descending, then TrackPercentile descending, then team name. Ranks are unique after this tie-break.",
        "incomplete": "Only teams with a submitted evaluation covering every active criterion are ranked. Pending teams are omitted unless explicitly requested.",
        "duplicates": "A unique index prevents two evaluations by the same judge for the same team. If multiple judges score one team, their totals are averaged so a second score cannot inflate rank.",
    }


def filter_leaderboard(
    results: dict,
    *,
    search: Optional[str] = None,
    domain: Optional[str] = None,
    track_id: Optional[str] = None,
    status: str = "completed",
    finalists: Optional[int] = None,
) -> list[dict]:
    rows = list(results["rows"])
    if status == "all":
        rows = rows + list(results["pending_rows"])
    elif status == "pending":
        rows = list(results["pending_rows"])
    if track_id:
        rows = [row for row in rows if row.get("track_id") == track_id]
    domain_q = (domain or "").strip().casefold()
    if domain_q:
        rows = [row for row in rows if (row.get("domain") or "").strip().casefold() == domain_q]
    query = (search or "").strip().casefold()
    if query:
        rows = [
            row
            for row in rows
            if query
            in " ".join(
                [
                    row.get("registration_id") or "",
                    row.get("reference_id") or "",
                    row.get("team_name") or "",
                    row.get("ps_id") or "",
                    row.get("problem_statement") or "",
                    row.get("theme") or "",
                    row.get("domain") or "",
                    row.get("track_name") or "",
                ]
            ).casefold()
        ]
    completed = [row for row in rows if row.get("evaluation_status") == "completed"]
    pending = [row for row in rows if row.get("evaluation_status") != "completed"]
    completed.sort(key=_sort_key)
    pending.sort(key=lambda row: ((row.get("track_name") or "").casefold(), (row.get("team_name") or "").casefold()))
    return [serialize_leaderboard_row(row, finalists) for row in completed + pending]


def _header_fill() -> PatternFill:
    return PatternFill("solid", fgColor="1F2937")


def _header_font() -> Font:
    return Font(bold=True, color="FFFFFF")


def _thin_border() -> Border:
    side = Side(style="thin", color="D1D5DB")
    return Border(top=side, bottom=side, left=side, right=side)


def _style_sheet(sheet: Worksheet, number_columns: Optional[Iterable[str]] = None, percent_columns: Optional[Iterable[str]] = None) -> None:
    border = _thin_border()
    for cell in sheet[1]:
        cell.font = _header_font()
        cell.fill = _header_fill()
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = border
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="center")
            cell.border = border
    sheet.freeze_panes = "A2"
    if sheet.max_row >= 1 and sheet.max_column >= 1:
        sheet.auto_filter.ref = sheet.dimensions
    for index, column in enumerate(sheet.columns, start=1):
        letter = get_column_letter(index)
        width = min(max((len(str(cell.value if cell.value is not None else "")) for cell in column), default=0) + 3, 42)
        sheet.column_dimensions[letter].width = max(width, 12)
    for letter in number_columns or []:
        for cell in sheet[letter][1:]:
            if isinstance(cell.value, (int, float)):
                cell.number_format = "0.00"
    for letter in percent_columns or []:
        for cell in sheet[letter][1:]:
            if isinstance(cell.value, (int, float)):
                cell.number_format = "0.00"


def build_results_workbook(results: dict, finalists: Optional[int] = None) -> bytes:
    workbook = Workbook()
    leaderboard = workbook.active
    leaderboard.title = "Final Leaderboard"
    leaderboard.append(
        [
            "Overall Rank",
            "Team Name",
            "Registration ID",
            "Problem Statement",
            "Track",
            "Raw Score",
            "Max Score",
            "Raw Percentage",
            "Track Percentile",
            "Final Score",
            "Track Rank",
            "Finalist",
        ]
    )
    for row in results["rows"]:
        payload = serialize_leaderboard_row(row, finalists)
        leaderboard.append(
            [
                payload["overall_rank"],
                payload["team_name"],
                payload["registration_id"],
                payload["problem_statement"],
                payload["track_name"],
                payload["raw_score"],
                payload["max_score"],
                payload["raw_percentage"],
                payload["track_percentile"],
                payload["final_score"],
                payload["track_rank"],
                "FINALIST" if payload["is_finalist"] else "",
            ]
        )
    _style_sheet(leaderboard, number_columns=("F", "G", "H", "I", "J"))

    stats_sheet = workbook.create_sheet("Track Statistics")
    stats_sheet.append(
        [
            "Track",
            "Total Teams",
            "Evaluated Teams",
            "Pending Teams",
            "Mean",
            "Std Dev",
            "Minimum",
            "Maximum",
            "Normalization Status",
        ]
    )
    for stat in results["track_stats"]:
        payload = serialize_track_stat(stat)
        stats_sheet.append(
            [
                payload["track_name"],
                payload["total_teams"],
                payload["evaluated_teams"],
                payload["pending_teams"],
                payload["mean_raw_score"],
                payload["stddev"],
                payload["minimum"],
                payload["maximum"],
                payload["normalization_status"],
            ]
        )
    _style_sheet(stats_sheet, number_columns=("E", "F", "G", "H"))

    detail_sheet = workbook.create_sheet("Detailed Evaluations")
    detail_sheet.append(["Team", "Registration ID", "Track", "Judge", "Criterion", "Score", "Max Marks"])
    criteria = results["criteria"]
    for row in results["rows"]:
        scores = row.get("criterion_scores") or {}
        judge_name = row.get("judge_name") or ""
        for criterion in criteria:
            detail_sheet.append(
                [
                    row.get("team_name", ""),
                    row.get("registration_id", ""),
                    row.get("track_name", ""),
                    judge_name,
                    criterion.get("name", ""),
                    round2(scores.get(criterion.get("id"))),
                    criterion.get("max_marks"),
                ]
            )
    _style_sheet(detail_sheet, number_columns=("F", "G"))

    method = workbook.create_sheet("Normalization Method")
    method_rows = [
        ["NMIET SIH Internal Hackathon — Evaluation Ranking"],
        [""],
        ["Raw Score", methodology_text()["raw_score"]],
        ["Raw Percentage", methodology_text()["raw_percentage"]],
        ["Track Percentile", methodology_text()["track_percentile"]],
        ["Final Score", methodology_text()["final_score"]],
        [""],
        ["Worked example"],
        ["RawPercentage = (RawScore / 70) × 100"],
        ["TrackPercentile = 100 × (teams with a lower raw score) / (n − 1)"],
        ["FinalScore = (RawPercentage × 0.70) + (TrackPercentile × 0.30)"],
        [""],
        ["Example: 68/70 with track percentile 100"],
        ["RawPercentage = 97.14"],
        ["FinalScore = (97.14 × 0.70) + (100 × 0.30) = 98.00"],
        [""],
        ["Tie handling", methodology_text()["ties"]],
        ["Incomplete evaluations", methodology_text()["incomplete"]],
        ["Duplicate evaluations", methodology_text()["duplicates"]],
        [""],
        ["Generated at (UTC)", datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M")],
    ]
    for row in method_rows:
        method.append(row)
    method["A1"].font = Font(bold=True, size=14)
    for cell in method["A"]:
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    for cell in method["B"]:
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    method.column_dimensions["A"].width = 28
    method.column_dimensions["B"].width = 88
    method.freeze_panes = "A2"

    output = BytesIO()
    workbook.save(output)
    return output.getvalue()
