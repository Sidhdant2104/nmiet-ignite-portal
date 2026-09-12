"""Tests for the registration Excel Statistics sheet."""

import unittest

from openpyxl import Workbook

from app.services.participation_statistics import (
    add_statistics_sheet,
    build_statistics_rows,
    classify_gender,
    student_identity,
)


def _person(name, email, gender="Male", year="Third Year", department="Computer Engineering", mobile="9000000001", roll="1"):
    return {
        "name": name,
        "email": email,
        "gender": gender,
        "year": year,
        "department": department,
        "mobile": mobile,
        "roll": roll,
    }


def _reg(registration_id, theme="Smart Automation", category="Software", status="PPT Submitted", members=None, leader=None):
    if members is None:
        members = [
            _person("A", f"{registration_id}-a@college.edu", "Female", "Second Year", "Information Technology (IT)", "9222222221"),
            _person("B", f"{registration_id}-b@college.edu", "Male", "Final Year", "Computer Engineering", "9222222222"),
        ]
    return {
        "registration_id": registration_id,
        "status": status,
        "isDeleted": False,
        "team": {"teamName": registration_id, "psId": "SIH26001", "psTitle": "Sample PS", "theme": theme, "category": category},
        "leader": leader or _person("Leader", f"{registration_id}-lead@college.edu", year="Third Year", mobile="9111111111"),
        "members": members,
    }


class GenderAndIdentityTests(unittest.TestCase):
    def test_gender_never_inferred_from_name(self):
        self.assertEqual(classify_gender(""), "Unknown")
        self.assertEqual(classify_gender(None), "Unknown")
        self.assertEqual(classify_gender("Male"), "Male")
        self.assertEqual(classify_gender("Female"), "Female")
        self.assertEqual(classify_gender("Priya"), "Unknown")

    def test_identity_prefers_email_then_phone(self):
        self.assertEqual(student_identity({"email": "A@X.com", "mobile": "900"}), ("email", "a@x.com"))
        self.assertEqual(student_identity({"email": "", "mobile": "900-111"}), ("phone", "900111"))
        self.assertIsNone(student_identity({"name": "Only Name", "roll": "15"}))


class StatisticsComputeTests(unittest.TestCase):
    def test_counts_unique_students_and_duplicate_email(self):
        shared = _person("Dup", "shared@college.edu", "Female", year="Second Year", mobile="9333333333")
        registrations = [
            _reg("T1", members=[shared, _person("C", "t1-c@college.edu", "Male", year="Final Year", mobile="9444444444")]),
            _reg("T2", members=[shared, _person("D", "t2-d@college.edu", "Female", year="Final Year", mobile="9555555555")]),
        ]
        stats = build_statistics_rows(registrations, tracks=[])
        self.assertEqual(stats["teams"], 2)
        self.assertEqual(stats["student_instances"], 6)
        self.assertEqual(stats["unique_students"], 5)
        self.assertEqual(stats["gender_counts"]["Female"], 2)
        self.assertEqual(stats["duplicate_email"], 1)
        self.assertEqual(stats["year_counts"]["Third Year"], 2)
        self.assertEqual(stats["year_counts"]["Final Year"], 2)
        self.assertEqual(stats["year_counts"]["Second Year"], 1)

    def test_missing_gender_is_unknown_not_guessed(self):
        registrations = [_reg("T1", leader=_person("X", "x@college.edu", gender=""), members=[])]
        stats = build_statistics_rows(registrations, tracks=[])
        self.assertEqual(stats["gender_counts"]["Unknown"], 1)
        self.assertEqual(stats["missing_gender"], 1)
        self.assertFalse(stats["gender_field_present"])

    def test_workbook_contains_statistics_sheet(self):
        workbook = Workbook()
        workbook.active.title = "Registrations"
        add_statistics_sheet(workbook, [_reg("T1")], tracks=[])
        self.assertIn("Statistics", workbook.sheetnames)
        sheet = workbook["Statistics"]
        values = [cell.value for row in sheet.iter_rows(max_col=1) for cell in row if cell.value]
        self.assertTrue(any("Total Teams Participated" in str(value) for value in values))
        self.assertTrue(any("Data quality" in str(value) for value in values))
        self.assertTrue(any("Year × Department" in str(value) for value in values))


if __name__ == "__main__":
    unittest.main()
